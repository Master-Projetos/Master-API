import logging
import re
from io import BytesIO

import openpyxl
import pandas as pd
import requests

from app.core.constants import SWAP_SHEET_URL
from app.modules.swap.common import (
    CHECKED,
    DATE_COLUMNS,
    MISSING_INFO,
    UNCHECKED,
    UNTITLED_SECTION,
    clean_text,
    parse_date,
    parse_number,
    parse_text,
    split_untitled,
)

logger = logging.getLogger(__name__)

SHEET_NAME = 'Andamento Documentação SWAP'

# Sections are a "SWAP" row followed by a title row. Which of them the API
# returns is set by SWAP_SECTIONS
SECTION_TITLE_PREFIX = 'SWAP'
SECTION_TITLES = {
    'FIBRA APAGADA | ENTRONCAMENTO | TRANSPORTE': UNTITLED_SECTION,
    'LAMBDA': 'lambdas',
    'COLOCATION': 'colocation',
    'CAPACIDADE': 'capacity',
    'LAN TO LAN': 'lan-to-lan',
    'OUTROS': 'others',
}

FETCH_TIMEOUT_SECONDS = 30

# Each contract has its own header, with shifting columns, so match by label
HEADER_LABELS = {'CESSÃO', 'CIRCUITO'}

# Sub-table headers put the partner name where this label would be
PARTNER_LABEL = 'PARCEIRA'

COLUMNS = {
    'CONTRATO / STATUS': 'contract',
    'OBS': 'notes',
    'VIGÊNCIA': 'term',
    'CESSÃO': 'assignment',
    'CIRCUITO': 'circuit',
    'ACORDO': 'agreement',
    'QNT DE FIBRAS': 'quantity',
    'QUANTIDADE': 'quantity',
    'DISTÂNCIA (KM)': 'distance_km',
    'KM FIBRA': 'fiber_km',
    'KM X FIBRA': 'fiber_km',
    'STATUS DO ACORDO': 'status',
    'DATA DE CELEBRAÇÃO': 'signed_at',
    'DATA DE TÉRMINO': 'ends_at',
}

# Not every contract cell is merged; some are just left blank below
CONTRACT_COLUMNS = [
    'partner',
    'contract',
    'notes',
    'term',
    'assignment',
    'signed_at',
    'ends_at',
]

TEXT_COLUMNS = (
    'partner',
    'contract',
    'notes',
    'term',
    'assignment',
    'circuit',
    'agreement',
    'quantity',
    'status',
)
NUMERIC_COLUMNS = ('distance_km', 'fiber_km')

RESPONSE_COLUMNS = [
    'section',
    'contract_id',
    'partner',
    'contract_url',
    'contract_status',
    'notes',
    'term',
    'assignment',
    'circuit',
    'agreement',
    'quantity',
    'distance_km',
    'fiber_km',
    'status',
    'signed_at',
    'ends_at',
    'check_status',
]

# The contract cell mixes a link with notes like "ADITIVO", and line breaks
# can split the link itself, so everything after "http" is the link
CONTRACT_URL_PATTERN = re.compile(r'^(.*?)\s*(https?://.*)$')

# The sheet marks how far each circuit was checked only by the fill color of
# its CIRCUITO cell: green, red, or white/none for not checked yet
CHECK_COLORS = {
    'FF00FF00': CHECKED,
    'FFFF0000': MISSING_INFO,
}
NO_COLORS = {None, 'FFFFFFFF', '00000000'}


def _fill_color(cell) -> str | None:
    """ARGB fill of a cell. Theme colors have no ARGB, so count as none."""
    if cell.fill.fill_type != 'solid':
        return None

    rgb = cell.fill.fgColor.rgb
    return rgb if isinstance(rgb, str) else None


def _fetch_sheet() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Cells as they are, to read the layout; with merged cells filled, since
    a merge only keeps its value in the first cell; and their fill colors."""
    logger.info('Fetching swap spreadsheet')
    response = requests.get(SWAP_SHEET_URL, timeout=FETCH_TIMEOUT_SECONDS)
    response.raise_for_status()

    # pandas drops the merges, so the sheet is read with openpyxl
    sheet = openpyxl.load_workbook(BytesIO(response.content))[SHEET_NAME]
    raw = pd.DataFrame(sheet.values).map(clean_text)

    # Only downwards: a merge across columns is one long text, like a note
    # spanning circuit, agreement and quantity, so it belongs to the first
    filled = raw.copy()
    for merge in sheet.merged_cells.ranges:
        top, left = merge.min_row - 1, merge.min_col - 1
        filled.iloc[top : merge.max_row, left] = raw.iat[top, left]

    colors = pd.DataFrame([
        [_fill_color(cell) for cell in row] for row in sheet.iter_rows()
    ])

    return raw, filled, colors


def _is_header(cells: pd.Series) -> bool:
    return HEADER_LABELS <= set(cells.dropna().astype(str))


def _first_cell_only(cells: pd.Series) -> str | None:
    """Value of rows with just the first column filled: section titles, or
    partners of sub-tables that left out their header labels."""
    filled = cells.dropna()
    if len(filled) == 1 and filled.index[0] == 0:
        return str(filled.iloc[0])

    return None


def _parse_rows(
    raw: pd.DataFrame, filled: pd.DataFrame, colors: pd.DataFrame
) -> pd.DataFrame:
    """Read row by row, since titles and headers change how the next rows
    are read. Values come from `filled`, so merged cells reach every row."""
    records = []
    section: str | None = UNTITLED_SECTION
    header: dict[int, str] | None = None
    header_partner = None
    contract = 0
    after_prefix = False

    for row, cells in raw.iterrows():
        if cells.isna().all():
            continue

        if (value := _first_cell_only(cells)) is not None:
            if value.upper() == SECTION_TITLE_PREFIX:
                after_prefix = True
                continue

            if after_prefix:
                after_prefix = False
                section = SECTION_TITLES.get(value.upper())
                if section is None:
                    logger.warning('Unknown swap sheet section "%s"', value)
                header = None
                continue

            # A sub-table with only its partner keeps the columns above
            header_partner = value
            contract += 1
            continue

        after_prefix = False

        if _is_header(cells):
            header = {
                position: COLUMNS[str(label).upper()]
                for position, label in cells.items()
                if str(label).upper() in COLUMNS
            }
            # Its partner is merged down over the rows below
            partner = cells.get(0)
            header_partner = None if partner == PARTNER_LABEL else partner
            contract += 1
            continue

        if section is None:
            continue

        if header is None:
            logger.warning('Swap sheet row %s has no header, skipping', row)
            continue

        # The unfilled partner cell tells where a merged contract starts
        partner = cells.get(0)
        if pd.notna(partner) and partner != header_partner:
            contract += 1

        values = filled.loc[row]
        partner = values.get(0)
        if pd.isna(partner):
            partner = header_partner

        record = {name: values[position] for position, name in header.items()}
        record.update(section=section, contract_id=contract, partner=partner)

        circuit_position = next(
            position for position, name in header.items() if name == 'circuit'
        )
        record['circuit_color'] = colors.at[row, circuit_position]
        records.append(record)

    # Keeps columns that no table in the sheet happens to have
    columns = [
        'section',
        'contract_id',
        'partner',
        *COLUMNS.values(),
        'circuit_color',
    ]
    return pd.DataFrame(records, columns=list(dict.fromkeys(columns)))


def _split_contract(contract: str | None) -> tuple[str | None, str | None]:
    """Link and notes of a contract cell."""
    if not isinstance(contract, str):
        return None, None

    if match := CONTRACT_URL_PATTERN.match(contract):
        notes, url = match.groups()
        return url.replace(' ', ''), notes or None

    return None, contract


def _check_statuses(colors: pd.Series) -> pd.Series:
    """Check status of each circuit, from the color of its cell."""
    # A new color means the sheet's convention changed, so it's flagged
    unknown = set(colors.dropna()) - set(CHECK_COLORS) - NO_COLORS
    if unknown:
        logger.warning(
            'Unknown swap circuit colors %s, read as unchecked',
            sorted(unknown),
        )

    return colors.map(CHECK_COLORS).fillna(UNCHECKED)


def load_sections(sections: list[str]) -> pd.DataFrame:
    """Rows of the given sections of the sheet, one per circuit."""
    raw, filled, colors = _fetch_sheet()
    df = split_untitled(_parse_rows(raw, filled, colors))

    df = df[df['section'].isin(sections)].reset_index(drop=True)
    # A renamed title would otherwise return an empty table
    missing = set(sections) - set(df['section'])
    if missing:
        raise ValueError(f'No {sorted(missing)} sections in the swap sheet')

    df[CONTRACT_COLUMNS] = df.groupby('contract_id')[CONTRACT_COLUMNS].ffill()

    for column in NUMERIC_COLUMNS:
        df[column] = df[column].map(parse_number).astype('Float64')

    for column in DATE_COLUMNS:
        df[column] = pd.to_datetime(df[column].map(parse_date))

    for column in TEXT_COLUMNS:
        df[column] = df[column].map(parse_text).astype('string')

    contracts = df['contract'].map(_split_contract)
    df['contract_url'] = contracts.str[0].astype('string')
    df['contract_status'] = contracts.str[1].astype('string')

    df['check_status'] = _check_statuses(df['circuit_color'])

    # Numbered from 1, and unique even when several sections are returned
    df['contract_id'] = pd.factorize(df['contract_id'])[0] + 1

    return df[RESPONSE_COLUMNS]
